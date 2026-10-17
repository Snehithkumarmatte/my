"""Data harvester tool.

Reads an invoice file (PDF, DOCX or scanned PNG), gets its raw text and
asks Bedrock to turn that text into structured invoice JSON.

   PDF   -> pdfplumber text layer, Tesseract OCR if the PDF is a scan
   DOCX  -> python-docx paragraphs and tables
   PNG   -> Tesseract OCR
"""

from __future__ import annotations

import json
import logging
import os
from pathlib import Path
from typing import Any

from langchain_core.tools import tool

from backend.llm import invoke_bedrock_chat
from backend.models import ExtractedInvoice, ExtractedLineItem, ExtractionResult
from backend.config.settings import BEDROCK_REASONING_MODEL
from backend.tools.rule_parser import parse_invoice_text

logger = logging.getLogger(__name__)

FORMATS = {".pdf": "pdf", ".docx": "docx", ".png": "image"}
OCR_LANGUAGES = os.getenv("TESSERACT_LANG", "eng+spa+deu")

EXTRACTION_PROMPT = """\
You are an expert invoice data extractor.  Your job is to extract
EVERY piece of information from the invoice and return it as JSON.

Return ONLY valid JSON — no markdown fences, no commentary.

The JSON schema you MUST follow:
{
 "invoice_no": "<string or null>",
 "invoice_date": "<string in any date format found, or null>",
 "po_number": "<Purchase Order reference like PO-1001, or null>",
 "vendor_id": "<Vendor ID like VEND-001, or null>",
 "vendor_name": "<string or null>",
 "currency": "<ISO 4217 code like USD, EUR, INR, GBP — or null>",
 "total_amount": <number or null>,
 "tax_amount": <number or null>,
 "subtotal": <number or null>,
 "line_items": [
   {
     "item_code": "<SKU code like SKU-001, or null>",
     "description": "<item description in ORIGINAL language — do NOT translate>",
     "qty": <number or null>,
     "unit_price": <number or null>,
     "total": <number or null>
   }
 ],
 "detected_language": "<ISO 639-1 code of the document language, e.g. en, de, es>",
 "additional_fields": { "<any other key-value pairs you find>" }
}

Rules:
- Extract ALL line items, even if the table has many rows.
- Preserve the original language of descriptions — do NOT translate.
- For currency, convert symbols to ISO codes ($ → USD, € → EUR, ₹ → INR, £ → GBP).
 CRITICAL: Set currency to null if NO currency symbol or ISO code is EXPLICITLY AND VISIBLY
 present in the document. NEVER guess, infer, or assume a currency from the country,
 company name, or amount magnitude. Only set it if you can directly see it in the text.
- If a field is not present in the document, set it to null.
- Numbers must be JSON numbers, not strings.
- Do NOT wrap the JSON in ```json``` markers.
"""


@tool
def data_harvester_tool(file_path: str, email_id: str = "unknown") -> dict:
   """Extract structured invoice data from a PDF, DOCX or PNG file.

   Gets the raw text from the file, then asks Bedrock to parse it into
   invoice fields and line items.

   Args:
       file_path: Absolute path to the invoice file.
       email_id: Email ID this invoice belongs to.

   Returns:
       ExtractionResult as a dict.
   """
   path = Path(file_path)
   result = {"email_id": email_id, "filename": path.name, "file_path": str(path)}

   if not path.exists():
       return ExtractionResult(
           **result, file_format="unknown", extraction_errors=[f"File not found: {path}"],
       ).model_dump()

   file_format = FORMATS.get(path.suffix.lower(), "unknown")
   raw_text, parsed, errors = "", {}, []

   try:
       if file_format == "unknown":
           errors.append(f"Unsupported format: {path.suffix.lower()}")
       else:
           logger.info("Reading %s: %s", file_format.upper(), path.name)
           raw_text = read_text(path, file_format)
           if raw_text.strip():
               parsed = parse_fields(raw_text, path.name)
           else:
               errors.append(f"No text extracted from {file_format.upper()}")
   except Exception as exc:
       errors.append(f"Extraction failed for {path.name}: {exc}")
       logger.error("Extraction error for %s: %s", path.name, exc)

   invoice = None
   if parsed:
       try:
           invoice = to_invoice(parsed)
       except Exception as exc:
           errors.append(f"Failed to build invoice model: {exc}")
           logger.error("Model build error: %s", exc)

   if invoice:
       logger.info(
           "%s -> Invoice#: %s | PO: %s | Vendor: %s | Items: %d | Total: %s %s",
           path.name, invoice.invoice_no or "N/A", invoice.po_number or "N/A",
           invoice.vendor_id or invoice.vendor_name or "N/A", len(invoice.line_items),
           invoice.currency or "?", invoice.total_amount,
       )

   return ExtractionResult(
       **result,
       file_format=file_format,
       raw_text=raw_text,
       extracted_invoice=invoice,
       extraction_errors=errors,
   ).model_dump()


# reading text from files

def read_text(path: Path, file_format: str) -> str:
   if file_format == "pdf":
       return read_pdf(path)
   if file_format == "docx":
       return read_docx(path)
   return ocr_image(path)


def read_pdf(path: Path) -> str:
   """Use the PDF's text layer (plus its tables). If there is no text,
   the PDF is a scan, so render the pages and OCR them."""
   import pdfplumber

   parts: list[str] = []
   with pdfplumber.open(path) as pdf:
       for page in pdf.pages:
           parts.append(page.extract_text() or "")
           for table in page.extract_tables() or []:
               for row in table:
                   if row:
                       parts.append(" | ".join(cell or "" for cell in row))

   text = "\n".join(parts).strip()
   if text:
       return text

   import pytesseract
   from pdf2image import convert_from_path

   pages = convert_from_path(str(path), dpi=250)
   return "\n".join(pytesseract.image_to_string(page, lang=OCR_LANGUAGES) for page in pages)


def read_docx(path: Path) -> str:
   from docx import Document

   doc = Document(str(path))
   parts = [p.text for p in doc.paragraphs if p.text.strip()]
   for table in doc.tables:
       for row in table.rows:
           parts.append(" | ".join(cell.text.strip() for cell in row.cells))
   return "\n".join(parts)


def ocr_image(path: Path) -> str:
   import pytesseract
   from PIL import Image

   return pytesseract.image_to_string(Image.open(path), lang=OCR_LANGUAGES)


# field parsing: Bedrock first, deterministic rules as the fallback

PARSER_MODE = os.getenv("EXTRACTION_PARSER", "auto").lower()   # auto | llm | rules


def parse_fields(raw_text: str, filename: str) -> dict:
   """Turn raw text into invoice fields.

   auto  -> Bedrock; if Bedrock fails or returns unusable JSON, the rule-based parser
   llm   -> Bedrock only (errors propagate)
   rules -> rule-based parser only (no model call at all)
   The raw text is already stored, so the choice never forces re-reading the file.
   """
   if PARSER_MODE == "rules":
       return parse_invoice_text(raw_text)
   try:
       logger.info("Parsing %s with Bedrock (%d chars) ...", filename, len(raw_text))
       parsed = parse_with_bedrock(raw_text)
       if not isinstance(parsed, dict) or not parsed:
           raise ValueError("empty or non-object JSON from model")
       parsed.setdefault("additional_fields", {})
       if isinstance(parsed["additional_fields"], dict):
           parsed["additional_fields"].setdefault("parser", "llm")
       return parsed
   except Exception as exc:
       if PARSER_MODE == "llm":
           raise
       logger.warning("Bedrock parsing failed for %s (%s) - using rule-based parser", filename, exc)
       return parse_invoice_text(raw_text)


# Bedrock parsing

def parse_with_bedrock(raw_text: str) -> dict:
   """Send the raw text to Bedrock and return the invoice JSON as a dict."""
   reply = invoke_bedrock_chat(
       model_id=BEDROCK_REASONING_MODEL,
       messages=[
           {"role": "system", "content": EXTRACTION_PROMPT},
           {"role": "user", "content": f"Extract all invoice data from the following text and return the JSON:\n\n{raw_text}"},
       ],
       max_tokens=4096,
       temperature=0.0,
   )
   return json.loads(strip_code_fence(reply or "{}"))


def strip_code_fence(text: str) -> str:
   """Remove ```json ... ``` if the model wrapped its answer anyway."""
   text = text.strip()
   if text.startswith("```"):
       text = text.split("\n", 1)[1] if "\n" in text else text[3:]
   if text.endswith("```"):
       text = text[:-3]
   return text.strip()


# dict -> models

def to_invoice(data: dict) -> ExtractedInvoice:
   lines = [
       ExtractedLineItem(
           item_code=item.get("item_code"),
           description=item.get("description"),
           qty=to_float(item.get("qty")),
           unit_price=to_float(item.get("unit_price")),
           total=to_float(item.get("total")),
       )
       for item in data.get("line_items") or []
       if isinstance(item, dict)
   ]
   return ExtractedInvoice(
       invoice_no=data.get("invoice_no"),
       invoice_date=data.get("invoice_date"),
       po_number=data.get("po_number"),
       vendor_id=data.get("vendor_id"),
       vendor_name=data.get("vendor_name"),
       currency=data.get("currency"),
       total_amount=to_float(data.get("total_amount")),
       tax_amount=to_float(data.get("tax_amount")),
       subtotal=to_float(data.get("subtotal")),
       line_items=lines,
       detected_language=data.get("detected_language"),
       additional_fields=data.get("additional_fields") or {},
   )


def to_float(value: Any) -> float | None:
   try:
       return None if value is None else float(value)
   except (ValueError, TypeError):
       return None